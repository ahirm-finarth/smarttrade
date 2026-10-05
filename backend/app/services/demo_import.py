"""Source-grounded, transactional import of synthetic reference data only."""

import csv
import hashlib
import json
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import Boolean, DateTime, Integer, Numeric, String, select
from sqlalchemy.orm import Session

from app.models.domain import (
    ApprovalEvent,
    CaseDocument,
    CaseParty,
    DemoRiskRule,
    DemoScreeningReference,
    Discrepancy,
    RiskEvent,
    TradeCase,
    TradeLine,
)

DATASET_KEY = "smart-trade-synthetic-v1"
SPECS = {
    "trade_cases": (TradeCase, "Trade Cases", ["case_id"]),
    "case_parties": (CaseParty, "Case Parties", ["case_id", "party_role", "party_name"]),
    "case_documents": (CaseDocument, "Case Documents", ["case_id", "document_id"]),
    "trade_lines": (TradeLine, "Trade Lines", ["case_id", "source", "item_no"]),
    "discrepancies": (Discrepancy, "Discrepancies", ["case_id", "finding_id"]),
    "risk_events": (RiskEvent, "Risk Events", ["case_id", "risk_id"]),
    "approval_events": (
        ApprovalEvent,
        "Approval Events",
        ["case_id", "event_time", "actor_id", "action"],
    ),
    "risk_rules": (DemoRiskRule, "Risk Rules", ["rule_id"]),
    "reference_screening": (DemoScreeningReference, "Demo Screening", ["reference_id"]),
}


class SourceError(ValueError):
    pass


def normalize(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value.lower() in {"nan", "null", "none"}:
            return None
    if isinstance(value, float) and value != value:
        return None
    return value


def read_sources(source: Path) -> dict[str, list[dict]]:
    """Accept a CSV directory or the supplied XLSX workbook, never mix sources."""
    bundle = {}
    if source.is_dir():
        for name in SPECS:
            path = source / f"{name}.csv"
            if path.exists():
                with path.open(encoding="utf-8-sig", newline="") as handle:
                    reader = csv.DictReader(handle)
                    validate_headers(name, reader.fieldnames or [])
                    bundle[name] = list(reader)
        if not bundle:
            workbooks = list(source.glob("*.xlsx"))
            if len(workbooks) == 1:
                return read_sources(workbooks[0])
    elif source.suffix.lower() == ".xlsx":
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            for name, (_, sheet, _) in SPECS.items():
                if sheet not in workbook.sheetnames:
                    continue
                rows = workbook[sheet].iter_rows(values_only=True)
                headers = list(next(rows, []))
                validate_headers(name, headers)
                bundle[name] = [dict(zip(headers, row, strict=True)) for row in rows if any(row)]
        finally:
            workbook.close()
    else:
        raise SourceError("Source must be a CSV directory or XLSX workbook")
    if not bundle.get("trade_cases"):
        raise SourceError("trade_cases source is missing or empty")
    return bundle


def validate_headers(name, headers):
    model, _, required = SPECS[name]
    allowed = set(model.__table__.columns.keys()) - {
        "id",
        "dataset_key",
        "case_pk",
        "source_key",
        "created_at",
        "updated_at",
    }
    if name not in {"trade_cases", "risk_rules", "reference_screening"}:
        allowed.add("case_id")
    if len(headers) != len(set(headers)):
        raise SourceError(f"{name}: duplicate column names")
    if set(required) - set(headers):
        raise SourceError(f"{name}: required columns missing")
    if set(headers) - allowed:
        raise SourceError(f"{name}: unrecognized columns; review the mapping")


def parse_value(value, column):
    value = normalize(value)
    if value is None:
        return None
    type_ = column.type
    if isinstance(type_, Boolean):
        if str(value).lower() in {"y", "true", "1"}:
            return True
        if str(value).lower() in {"n", "false", "0"}:
            return False
        raise ValueError("Invalid boolean")
    if isinstance(type_, Numeric):
        parsed = Decimal(str(value))
        quantum = Decimal(1).scaleb(-type_.scale)
        if not parsed.is_finite() or parsed != parsed.quantize(quantum):
            raise ValueError("Invalid numeric precision")
        if abs(parsed) >= Decimal(10) ** (type_.precision - type_.scale):
            raise ValueError("Numeric overflow")
        return parsed
    if isinstance(type_, Integer):
        parsed = Decimal(str(value))
        if parsed != parsed.to_integral_value():
            raise ValueError("Invalid integer")
        return int(parsed)
    if isinstance(type_, DateTime) or isinstance(getattr(type_, "impl", None), DateTime):
        parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
        if parsed.tzinfo is None:
            raise ValueError("Timestamp timezone required")
        return parsed.astimezone(UTC)
    parsed = str(value)
    if isinstance(type_, String) and type_.length and len(parsed) > type_.length:
        raise ValueError("String too long")
    return parsed


def prepare(bundle):
    prepared = {}
    for name, rows in bundle.items():
        model, _, required = SPECS[name]
        prepared[name] = []
        keys = set()
        for index, raw in enumerate(rows, 2):
            validate_headers(name, list(raw))
            try:
                row = {
                    key: parse_value(value, model.__table__.columns[key])
                    if key != "case_id" or name == "trade_cases"
                    else normalize(value)
                    for key, value in raw.items()
                }
                if any(row[key] is None for key in required):
                    raise ValueError("Required identifier is empty")
                key = tuple(str(row[k]) for k in required)
                if key in keys:
                    raise ValueError("Duplicate source identifier")
                keys.add(key)
                prepared[name].append(row)
            except (ValueError, TypeError, InvalidOperation, OverflowError):
                raise SourceError(f"{name}: invalid data on row {index}") from None
    cases = {row["case_id"] for row in prepared["trade_cases"]}
    for name, rows in prepared.items():
        if name not in {"trade_cases", "risk_rules", "reference_screening"}:
            if any(row["case_id"] not in cases for row in rows):
                raise SourceError(f"{name}: orphan case relationship")
    return prepared


def assign(entity, values):
    for key, value in values.items():
        if getattr(entity, key) != value:
            setattr(entity, key, value)


def ingest(session: Session, bundle: dict) -> dict[str, int]:
    """Caller owns the transaction. Only importer-owned records are reconciled."""
    prepared = prepare(bundle)
    cases = {}
    for row in prepared["trade_cases"]:
        case = session.scalar(select(TradeCase).where(TradeCase.case_id == row["case_id"]))
        if case and case.dataset_key != DATASET_KEY:
            raise SourceError("Demo case identifier conflicts with a non-demo record")
        if case is None:
            case = TradeCase(case_id=row["case_id"], dataset_key=DATASET_KEY)
            session.add(case)
        assign(case, row)
        cases[row["case_id"]] = case
    session.flush()
    for name, rows in prepared.items():
        model, _, identity = SPECS[name]
        if name == "trade_cases":
            continue
        if name in {"risk_rules", "reference_screening"}:
            id_field = identity[0]
            for row in rows:
                entity = session.scalar(
                    select(model).where(getattr(model, id_field) == row[id_field])
                )
                if entity and entity.dataset_key != DATASET_KEY:
                    raise SourceError("Demo reference conflicts with a non-demo record")
                if entity is None:
                    entity = model(**{id_field: row[id_field], "dataset_key": DATASET_KEY})
                    session.add(entity)
                assign(entity, row)
            continue
        for case_id, case in cases.items():
            existing = {
                entity.source_key: entity
                for entity in session.scalars(
                    select(model).where(
                        model.case_pk == case.id, model.source_key.startswith("demo:")
                    )
                )
            }
            wanted = set()
            for row in rows:
                if row["case_id"] != case_id:
                    continue
                encoded = json.dumps([str(row[key]) for key in identity], ensure_ascii=False)
                source_key = "demo:" + hashlib.sha256(encoded.encode()).hexdigest()[:59]
                wanted.add(source_key)
                entity = existing.get(source_key)
                if entity is None:
                    entity = model(case_pk=case.id, source_key=source_key)
                    session.add(entity)
                assign(entity, {key: value for key, value in row.items() if key != "case_id"})
            for key in existing.keys() - wanted:
                session.delete(existing[key])
    session.flush()
    return {name: len(rows) for name, rows in prepared.items()}
